"""
Export-layer tests (Gaps G4/G14 + the signed-download sweep finding):
drive the founder journey to Stage 3, then export every artefact in every
format and download one file through the signed /files route.
"""
from django.test import TestCase
from rest_framework.test import APIClient

from fundos.core.auth import issue_tokens
from tests.conftest_helpers import make_world


class ExportTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()

    def setUp(self):
        from django.core.management import call_command
        call_command("seed_initial_data", "--demo", verbosity=0)
        self.world = make_world()
        self.founder = self.world["a"]["founder"]
        self.deal = self.world["a"]["deal"]
        self.client_api = APIClient()
        access, _ = issue_tokens(self.founder)
        self.client_api.credentials(HTTP_AUTHORIZATION=f"Bearer {access}")
        self.base = f"/api/v1/deals/{self.deal.id}"
        self._drive_to_stage3()

    def _drive_to_stage3(self):
        c, base = self.client_api, self.base
        c.patch(f"{base}/ckb", {"fields": [
            {"fieldKey": "sector", "value": "SaaS"},
            {"fieldKey": "arr", "value": 120000000, "ccy": "INR"},
            {"fieldKey": "revenue", "value": 120000000, "ccy": "INR"},
            {"fieldKey": "burn", "value": 4000000, "ccy": "INR"},
        ]}, format="json")
        c.post(f"{base}/readiness/generate", {}, format="json")
        c.put(f"{base}/strategy/objectives",
              {"purposes": ["growth"], "timeline": "6m", "runway": "18m"},
              format="json")
        for what in ("peers", "raise", "valuation", "blueprint"):
            c.post(f"{base}/strategy/{what}/generate", {}, format="json")
        c.post(f"{base}/strategy/approve", {}, format="json")
        c.post(f"{base}/story/generate", {}, format="json")
        c.post(f"{base}/teaser/generate", {}, format="json")
        c.post(f"{base}/deck/outline", {}, format="json")
        c.post(f"{base}/deck/approve-outline", {}, format="json")
        c.post(f"{base}/deck/slides", {}, format="json")
        c.post(f"{base}/model/generate", {}, format="json")
        c.post(f"{base}/im/generate", {}, format="json")
        c.post(f"{base}/review/run", {}, format="json")

    def _assert_export(self, response, expected_fmt):
        self.assertEqual(response.status_code, 200, response.content[:300])
        body = response.json()
        self.assertIn("downloadUrl", body)
        self.assertEqual(body["format"], expected_fmt)
        return body

    def test_teaser_exports(self):
        for fmt in ("pdf", "pptx", "docx"):
            response = self.client_api.post(f"{self.base}/teaser/export",
                                            {"format": fmt}, format="json")
            self._assert_export(response, fmt)

    def test_deck_exports(self):
        for fmt, expected in (("pptx", "pptx"), ("pdf", "pdf"),
                              ("notes", "pdf")):
            response = self.client_api.post(f"{self.base}/deck/export",
                                            {"format": fmt}, format="json")
            self._assert_export(response, expected)

    def test_model_export_xlsx_has_live_formulas(self):
        response = self.client_api.post(f"{self.base}/model/export",
                                        {"format": "xlsx"}, format="json")
        body = self._assert_export(response, "xlsx")
        # Read back the workbook from local storage and check formulas.
        import os

        from django.conf import settings
        from openpyxl import load_workbook
        path = os.path.join(settings.FUNDOS_STORAGE_LOCAL_ROOT,
                            body["storageUri"])
        wb = load_workbook(path)
        formulas = [
            cell.value
            for name in wb.sheetnames if name != "Assumptions"
            for row in wb[name].iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=SUM(")
        ]
        self.assertTrue(formulas, "expected live =SUM() formulas in xlsx")

    def test_im_exports_including_watermark(self):
        for fmt, expected in (("pdf", "pdf"), ("docx", "docx"),
                              ("watermarked", "pdf")):
            response = self.client_api.post(f"{self.base}/im/export",
                                            {"format": fmt}, format="json")
            self._assert_export(response, expected)

    def test_reports(self):
        for route in ("readiness/report", "strategy/report",
                      "review/report"):
            response = self.client_api.get(f"{self.base}/{route}")
            self._assert_export(response, "pdf")

    def test_signed_download_roundtrip(self):
        """Sweep finding: signed URLs previously pointed at a route that
        did not exist. The full loop must work — and tampering must 403."""
        response = self.client_api.post(f"{self.base}/teaser/export",
                                        {"format": "pdf"}, format="json")
        # SAY WHY, don't raise KeyError. Reading the key straight out turned
        # a storage failure -- which returns 422 and an explanation -- into
        # "KeyError: 'downloadUrl'", which names nothing and sent a reader
        # looking for a contract change that had not happened.
        self.assertEqual(response.status_code, 200,
                         f"export failed: {response.json()}")
        url = response.json()["downloadUrl"]
        download = APIClient().get(url)          # signature IS the auth
        self.assertEqual(download.status_code, 200)
        content = b"".join(download.streaming_content)
        self.assertTrue(content.startswith(b"%PDF"))
        tampered = url.replace("sig=", "sig=00")
        self.assertEqual(APIClient().get(tampered).status_code, 403)

    def test_invalid_format_rejected(self):
        response = self.client_api.post(f"{self.base}/teaser/export",
                                        {"format": "exe"}, format="json")
        self.assertEqual(response.status_code, 422)
