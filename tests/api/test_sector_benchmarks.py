"""v22 — sector benchmark ingestion (Phase 2) and the Phase 1 fixes."""
import datetime as dt
from decimal import Decimal

from django.test import TestCase


def _write_workbook(path):
    """A miniature of the real sheet: three blocks, side by side and below.

    The layout is the point. Reading the left column as one table files the
    clubbed sub-sector groups as sectors, and bounding the raw table to the
    sector block loses most of the mappings.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sector Deal Data"
    ws.append(["Sector & Sub-Sector Deal Data"])
    ws.append([])
    ws.append(["1. Sector Attractiveness", None, None, None, None, None,
               None, None, None, None, "Sub-Sector source detail"])
    ws.append([])
    ws.append(["Sr. No.", "Name", "# Deals", "Raised ($Mn)",
               "Avg Ticket ($Mn)", "# Investors", "Velocity Score",
               "Ticket Score", "Investors Score", "Sample Quality",
               "Raw Sub-Sector", "# Deals", "Raised ($Mn)", "# Investors",
               "Clubbed Group"])
    ws.append([1, "Ecommerce", 606, 4977.55, 8.21, 193, 10, 2.5, 10, "OK",
               "D2C", 481, 2769.74, 139, "D2C & Consumer Brands"])
    ws.append([2, "Fintech", 373, 7387.05, 19.80, 128, 9.37, 10, 9.37, "OK",
               "Lending Tech", 165, 3823.12, 64, "Lending & Credit"])
    ws.append([3, "Agritech", 3, 12.0, 4.0, 2, None, None, None, None,
               "Gaming", 39, 318.31, 8, "Gaming & AR/VR"])
    # The raw table runs TALLER than the sector block beside it.
    ws.append([None, None, None, None, None, None, None, None, None, None,
               "Home Decor", 1, 0.59, 0, "D2C & Consumer Brands"])
    ws.append(["Total — Sector Table", None, 982, 12376.6, None, 323])
    ws.append([])
    ws.append(["2. Sub-Sector Attractiveness"])
    ws.append(["Sr. No.", "Name", "# Deals", "Raised ($Mn)",
               "Avg Ticket ($Mn)", "# Investors", "Velocity Score",
               "Ticket Score", "Investors Score", "Sample Quality"])
    ws.append([1, "D2C & Consumer Brands", 490, 2799.30, 5.71, 140, 10, 4, 10,
               "OK"])
    ws.append([2, "Lending & Credit", 168, 3830.31, 22.79, 64, 9.35, 10, 9.35,
               "OK"])
    ws.append([])
    # The sheet's own integrity checks sit below, carrying numbers in the
    # same columns.
    ws.append(["Deals in source extract (input)", None, 2699])
    wb.save(path)


class ParseWorkbookTests(TestCase):
    def setUp(self):
        import tempfile
        self.path = tempfile.mktemp(suffix=".xlsx")
        _write_workbook(self.path)

    def tearDown(self):
        import os
        try:
            os.unlink(self.path)
        except OSError:
            pass

    def _parse(self):
        from fundos.assessment.benchmarks import parse_workbook
        return parse_workbook(self.path)

    def test_the_three_blocks_are_read_separately(self):
        sectors, subs, report = self._parse()
        self.assertEqual([r["name"] for r in sectors],
                         ["Ecommerce", "Fintech", "Agritech"])
        self.assertEqual([r["name"] for r in subs],
                         ["D2C & Consumer Brands", "Lending & Credit"])

    def test_a_totals_row_is_not_imported_as_a_sector(self):
        sectors, _, _ = self._parse()
        self.assertNotIn("Total — Sector Table",
                         [r["name"] for r in sectors])

    def test_integrity_check_rows_below_the_table_are_excluded(self):
        """They carry a deal count in the same column and would be scored
        against as a sub-sector named 'Deals in source extract (input)'."""
        _, subs, _ = self._parse()
        for row in subs:
            self.assertNotIn("source extract", row["name"])

    def test_the_raw_mapping_table_is_read_past_the_sector_block(self):
        """It is taller than the block beside it — 104 raw labels against 18
        sectors in the real workbook. Bounding it to the block kept 21."""
        _, _, report = self._parse()
        pairs = dict(report["mappings"])
        self.assertEqual(pairs["D2C"], "D2C & Consumer Brands")
        self.assertEqual(pairs["Home Decor"], "D2C & Consumer Brands")

    def test_the_workbooks_own_scores_are_preserved(self):
        """An import reproduces the sheet rather than quietly recomputing
        it, so a figure an analyst can see is the figure that scores."""
        sectors, _, _ = self._parse()
        ecommerce = next(r for r in sectors if r["name"] == "Ecommerce")
        self.assertEqual(ecommerce["velocity_score"], Decimal("10"))
        self.assertEqual(ecommerce["ticket_score"], Decimal("2.5"))

    def test_a_thin_sample_is_flagged(self):
        sectors, _, _ = self._parse()
        agritech = next(r for r in sectors if r["name"] == "Agritech")
        self.assertEqual(agritech["sample_quality"], "Thin")


class ApplyRowsTests(TestCase):
    def setUp(self):
        import tempfile
        self.path = tempfile.mktemp(suffix=".xlsx")
        _write_workbook(self.path)

    def tearDown(self):
        import os
        try:
            os.unlink(self.path)
        except OSError:
            pass

    def _import(self, dry_run=False):
        from fundos.assessment.benchmarks import apply_rows, parse_workbook

        sectors, subs, report = parse_workbook(self.path)
        return apply_rows(sectors + subs, as_of_date=dt.date(2026, 8, 1),
                          source_file="test.xlsx", dry_run=dry_run,
                          mappings=report.get("mappings"))

    def test_a_dry_run_writes_nothing(self):
        from fundos.assessment.models import SectorDealData, SectorMapping

        result = self._import(dry_run=True)
        self.assertEqual(len(result["created"]), 5)
        self.assertEqual(SectorDealData.objects.count(), 0)
        self.assertEqual(SectorMapping.objects.count(), 0)

    def test_import_creates_rows_and_mappings(self):
        from fundos.assessment.models import SectorDealData, SectorMapping

        self._import()
        self.assertEqual(
            SectorDealData.objects.filter(level="sector").count(), 3)
        self.assertEqual(
            SectorDealData.objects.filter(level="sub_sector").count(), 2)
        self.assertEqual(SectorMapping.objects.count(), 4)
        self.assertEqual(
            SectorMapping.objects.get(raw_label="Home Decor").clubbed_group,
            "D2C & Consumer Brands")

    def test_reimport_updates_in_place_and_reports_the_move(self):
        """A refresh is MEANT to move ratings — newer market data should
        change what a fast-moving sector is. The dry run is the safeguard,
        so it has to name what would move."""
        from openpyxl import load_workbook

        from fundos.assessment.models import SectorDealData

        self._import()
        wb = load_workbook(self.path)
        wb["Sector Deal Data"]["C6"] = 900          # Ecommerce deal count
        wb.save(self.path)

        preview = self._import(dry_run=True)
        self.assertTrue(any("Ecommerce" in line and "606 -> 900" in line
                            for line in preview["updated"]), preview["updated"])
        self.assertEqual(
            SectorDealData.objects.get(level="sector",
                                       name="Ecommerce").deal_count, 606)

        applied = self._import()
        self.assertEqual(len(applied["created"]), 0)
        self.assertEqual(
            SectorDealData.objects.get(level="sector",
                                       name="Ecommerce").deal_count, 900)
        self.assertEqual(SectorDealData.objects.filter(level="sector").count(),
                         3, "a re-import must update, not duplicate")

    def test_a_changed_mapping_is_reported_before_it_is_applied(self):
        """Re-pointing a label changes which benchmark group a company is
        judged against — a category worth 20% of the rating."""
        from fundos.assessment.models import SectorMapping

        self._import()
        SectorMapping.objects.filter(raw_label="D2C").update(
            clubbed_group="Something Else")
        preview = self._import(dry_run=True)
        self.assertTrue(any("D2C" in line
                            for line in preview.get("mapping_changes", [])))


class PhaseOneFixTests(TestCase):
    """The three defects the medibuddy run exposed."""

    def test_a_revenue_model_with_no_percentages_seeds_as_a_draft(self):
        """v27.3 — superseded. These rows are KEPT.

        The model described the revenue streams instead of splitting them,
        which is the correct answer for most companies: they publish their
        streams and not the split. v27.2 stopped failing that as "Shares must
        total 100% (currently 0%)", but still discarded every row — so both
        17 Aug runs paid for the call and threw away four correct stream
        names, leaving the section blank with nothing on screen to explain
        why. A named stream with no percentage is still information, and the
        split is the one part of this the user actually knows.
        """
        from fundos.profile.services import _records_items_to_api

        rows, dropped = _records_items_to_api(
            "revenue_model",
            [{"label": "revenue model description"}, {"label": "subscription"}])
        self.assertEqual([r["label"] for r in rows],
                         ["revenue model description", "subscription"])
        self.assertNotIn("pct", rows[0])
        self.assertIn("draft", dropped[0])

    def test_a_partial_split_keeps_the_priced_rows(self):
        from fundos.profile.services import _records_items_to_api

        rows, dropped = _records_items_to_api(
            "revenue_model", [{"label": "A", "pct": 60}, {"label": "B"}])
        self.assertEqual([r["label"] for r in rows], ["A"])
        self.assertTrue(dropped)

    def test_fan_out_roles_do_not_resend_the_whole_corpus(self):
        """A live run sent 24,081 characters to each of ten follow-up calls,
        re-reading a corpus the deep extract had already structured."""
        from fundos.llm.adapter import _trim_bulk_sources

        corpus = "x" * 31000
        trimmed = _trim_bulk_sources("company_profile_records",
                                     {"website_extract": corpus})
        self.assertLess(len(trimmed["website_extract"]), 7000)
        self.assertIn("omitted", trimmed["website_extract"])

    def test_the_deep_extract_still_sees_everything(self):
        from fundos.llm.adapter import _trim_bulk_sources

        corpus = "x" * 31000
        kept = _trim_bulk_sources("company_profile_deep_extract",
                                  {"website_extract": corpus})
        self.assertEqual(kept["website_extract"], corpus)
