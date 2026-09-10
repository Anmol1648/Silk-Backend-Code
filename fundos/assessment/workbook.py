"""Read the assessment workbook as CONFIGURATION, not just as benchmark data.

WHY THE WHOLE WORKBOOK, NOT ONE SHEET
--------------------------------------
The workbook is not a data file with a benchmark tab. It is the scoring engine
expressed as formulae: the 80-row input contract, 52 stage-dependent rubrics,
24 written band definitions, the seven-category node tree with its weights,
the scalar constants on the Deal Scorecard panel, and — most useful of all —
a Data Dictionary telling an analyst what each parameter means, where to find
it and what to do when it is absent.

Until v24 all of that was derived TWICE: once by whoever transcribed the spec
into `spec_config_data.py`, and once by the seeder's own assumptions about how
the categories fit together. Both copies were free to drift from the sheet
that is actually authoritative, and both did:

  * sub-sector keys were written `SUB_TAM` where the workbook says
    `SEC_TAM_SUB`, and the suffix is load-bearing — the sheet's own band
    formula does SUBSTITUTE(key,"_SUB","") to find the rubric;
  * six percentile parameters were invented that the workbook does not have,
    inflating an input contract the sheet counts as exactly 80;
  * those percentiles were scored as their own value, where the workbook
    bands them at 8/6/4 and scores 9/7/5/3 like everything else;
  * rows under five deals were scored, where the workbook excludes them.

None of those were reasoning failures about the domain. They were all the same
structural mistake — a second copy of the truth. Reading the sheet directly is
the fix, and a workbook revision becomes a re-import rather than a release.

WHAT THIS MODULE DOES NOT DO
----------------------------
It does not evaluate formulae. Cut-points, weights, band definitions and
constants are VALUES in the sheet and are read as values. The arithmetic those
formulae describe is implemented in `services.py`, and the acceptance test for
this whole exercise is that a filled workbook and our engine produce the same
scores from the same inputs.
"""
import logging
import re

logger = logging.getLogger(__name__)

INPUT_SHEET = "Company Profile"
RUBRIC_SHEET = "Rubrics"
ANCHOR_SHEET = "Qualitative Anchors"
DICT_SHEET = "Data Dictionary"
SCORECARD_SHEET = "Deal Scorecard"

# The seven category sheets, in the order the workbook lists them.
CATEGORY_SHEETS = [
    ("A", "A. Team"), ("B", "B. Financials"), ("C", "C. Business Quality"),
    ("D", "D. Deal Dynamics"), ("E", "E. Sector"), ("F", "F. Sub-Sector"),
    ("G", "G. Mandate Context"),
]

STAGES = ["Seed", "Series A", "Series B", "Growth"]

# Rubrics Table 1: the four stage column-triples, then the resolved ACTIVE
# STAGE triple which we deliberately ignore — it mirrors whichever stage the
# sheet happens to be set to, and we need all four.
STAGE_COLUMNS = {"Seed": (7, 8, 9), "Series A": (10, 11, 12),
                 "Series B": (13, 14, 15), "Growth": (16, 17, 18)}
# Table 2 (range): ideal min/max per stage, then shared tolerances.
RANGE_COLUMNS = {"Seed": (7, 8), "Series A": (9, 10), "Series B": (11, 12),
                 "Growth": (13, 14)}
RANGE_TOL_GOOD, RANGE_TOL_FAIR = 15, 16


class WorkbookError(ValueError):
    """A structural problem worth refusing to import over."""


def _cell(ws, row, col):
    v = ws.cell(row, col).value
    return v


def _text(ws, row, col):
    v = _cell(ws, row, col)
    return "" if v is None else str(v).strip()


def _num(value):
    from decimal import Decimal, InvalidOperation
    if value in (None, "", "-"):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = re.sub(r"[^\d.\-]", "", str(value))
    if not text or text in ("-", ".", "-."):
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


# ---------------------------------------------------------------------------
# 1. The input contract — Company Profile rows 22..101
# ---------------------------------------------------------------------------

def read_inputs(wb):
    """The 80-row contract: ref, input key, name, unit.

    This is the interface between step 1 and step 2, and its exact key
    spelling matters more than anything else in the file.
    """
    ws = wb[INPUT_SHEET]
    rows, seen = [], set()
    for r in range(22, 102):
        key = _text(ws, r, 3)
        if not key:
            continue
        if key in seen:
            raise WorkbookError(
                f"{INPUT_SHEET} row {r}: input key {key!r} appears twice. "
                f"Every category sheet resolves by key, so a duplicate makes "
                f"the lookup ambiguous.")
        seen.add(key)
        rows.append({
            "ref_code": _text(ws, r, 2),
            "input_key": key,
            "name": _text(ws, r, 4),
            "unit": _text(ws, r, 5),
            "row": r,
        })
    if not rows:
        raise WorkbookError(
            f"{INPUT_SHEET} rows 22-101 carried no input keys. Column C is "
            f"the key column; check the sheet has not been restructured.")
    return rows


# ---------------------------------------------------------------------------
# 2. Rubrics — 49 monotonic + 3 range, each across four stages
# ---------------------------------------------------------------------------

def read_rubrics(wb):
    ws = wb[RUBRIC_SHEET]
    monotonic, ranges = [], []

    for r in range(13, 62):
        key = _text(ws, r, 2)
        if not key:
            continue
        row = {"input_key": key, "metric_name": _text(ws, r, 3),
               "unit": _text(ws, r, 4), "direction": _text(ws, r, 5),
               "ref_code": _text(ws, r, 6),
               "rationale": _text(ws, r, 22), "stage_cuts": {}}
        for stage, (a, b, c) in STAGE_COLUMNS.items():
            row["stage_cuts"][stage] = (_num(_cell(ws, r, a)),
                                        _num(_cell(ws, r, b)),
                                        _num(_cell(ws, r, c)))
        monotonic.append(row)

    for r in range(67, 70):
        key = _text(ws, r, 2)
        if not key:
            continue
        row = {"input_key": key, "metric_name": _text(ws, r, 3),
               "unit": _text(ws, r, 4), "direction": "Range",
               "ref_code": _text(ws, r, 6),
               "rationale": _text(ws, r, 22),
               "good_tol": _num(_cell(ws, r, RANGE_TOL_GOOD)),
               "fair_tol": _num(_cell(ws, r, RANGE_TOL_FAIR)),
               "stage_bands": {}}
        for stage, (lo, hi) in RANGE_COLUMNS.items():
            row["stage_bands"][stage] = (_num(_cell(ws, r, lo)),
                                         _num(_cell(ws, r, hi)))
        ranges.append(row)

    if not monotonic:
        raise WorkbookError(
            f"{RUBRIC_SHEET} rows 13-61 carried no rubric keys — the sheet "
            f"layout has changed and no numeric parameter could be banded.")
    return monotonic, ranges


# ---------------------------------------------------------------------------
# 3. Qualitative anchors — the four written band definitions
# ---------------------------------------------------------------------------

def read_anchors(wb):
    ws = wb[ANCHOR_SHEET]
    out = []
    for r in range(6, 40):
        key = _text(ws, r, 2)
        if not key:
            continue
        out.append({
            "input_key": key, "ref_code": _text(ws, r, 3),
            "parameter_name": _text(ws, r, 4),
            "scoring_basis": _text(ws, r, 5),
            "excellent_def": _text(ws, r, 6), "good_def": _text(ws, r, 7),
            "fair_def": _text(ws, r, 8), "poor_def": _text(ws, r, 9),
            "evidence_required": _text(ws, r, 10),
        })
    return out


# ---------------------------------------------------------------------------
# 4. The node tree — read from the seven category sheets
# ---------------------------------------------------------------------------

_BENCH_COLS = {"H": "velocity", "I": "ticket", "J": "investors"}


def _benchmark_from_formula(formula):
    """Which percentile column a node reads, from its own Value formula.

    Detected from the formula rather than from the row label because the label
    is prose and the column reference is the thing that actually decides which
    number is scored.
    """
    text = str(formula or "")
    if "Sector Deal Data" not in text:
        return "", ""
    col = re.search(r"'Sector Deal Data'!\$([HIJ])\$", text)
    metric = _BENCH_COLS.get(col.group(1)) if col else ""
    # Sectors occupy rows 6-23, clubbed sub-sectors rows 28-60.
    level = "sub_sector" if re.search(r"\$C\$2[89]|\$C\$[3-6]\d", text) \
        else "sector"
    return level, (metric or "")


def read_tree(wb, formulas=None):
    """Category → sub-item → child, with weights and the key feeding each leaf.

    The sheets share one layout: column B carries the level marker (a capital
    letter for the category, a digit for a sub-item, a lower-case letter for a
    child), C the label, D the weight, N the input key where the row is a leaf.
    A leaf with no input key is a node the workbook COMPUTES — the deal-database
    percentiles — and is identified from its Value formula.
    """
    categories, subitems, children, leaves = [], [], [], []

    for code, sheet_name in CATEGORY_SHEETS:
        if sheet_name not in wb.sheetnames:
            raise WorkbookError(f"the workbook has no sheet {sheet_name!r}.")
        ws = wb[sheet_name]
        # Percentile nodes are identified from their Value FORMULA, which a
        # values-only read cannot see — on a blank template the cell is empty,
        # so detection silently found nothing and category E and F lost their
        # three table-driven nodes each.
        fs = formulas[sheet_name] if formulas is not None else None
        current_sub = None
        for r in range(1, ws.max_row + 1):
            marker = _text(ws, r, 2)
            if not marker:
                continue
            label = _text(ws, r, 3)
            weight = _num(_cell(ws, r, 4))
            key = _text(ws, r, 14)

            if marker == code:                       # the category row
                if weight is not None:
                    categories.append({"code": code, "label": label,
                                       "weight": weight * 100})
                continue
            if marker.isdigit():                     # a sub-item
                current_sub = f"{code}.{marker}"
                if weight is None:
                    continue
                subitems.append({"code": current_sub, "parent_code": code,
                                 "label": label, "weight": weight * 100})
                if key:
                    leaves.append({"node": current_sub, "input_key": key})
                else:
                    level, metric = _benchmark_from_formula(
                        _cell(fs, r, 15) if fs is not None else None)
                    if metric:
                        leaves.append({"node": current_sub, "input_key": "",
                                       "benchmark_level": level,
                                       "benchmark_metric": metric,
                                       "label": label, "ref_code": current_sub,
                                       "category_code": code})
                continue
            if marker.isalpha() and marker.islower() and current_sub:
                child = f"{current_sub}.{marker}"
                if weight is None:
                    continue
                children.append({"code": child, "parent_code": current_sub,
                                 "label": label, "weight": weight * 100})
                if key:
                    leaves.append({"node": child, "input_key": key})

    if len(categories) != 7:
        raise WorkbookError(
            f"expected seven category rows across the category sheets, "
            f"found {len(categories)}. Every score depends on the seven "
            f"weights summing to 100.")
    total = sum(c["weight"] for c in categories)
    if abs(total - 100) > 1:
        raise WorkbookError(
            f"the seven category weights sum to {total}, not 100. Importing "
            f"this would make every score built on it wrong.")
    return categories, subitems, children, leaves


# ---------------------------------------------------------------------------
# 5. Scalar constants — the Deal Scorecard panel
# ---------------------------------------------------------------------------

CONSTANT_LABELS = {
    "score if excellent": "score_excellent",
    "score if good": "score_good",
    "score if fair": "score_fair",
    "score if poor": "score_poor",
    "excellent if percentile score ≥": "percentile_excellent",
    "good if percentile score ≥": "percentile_good",
    "fair if percentile score ≥": "percentile_fair",
    "minimum deals for a reliable percentile": "min_deals_for_percentile",
}
RATING_LABELS = ("excellent", "very good", "good", "average", "challenging")


def read_constants(wb):
    """Band scores, percentile cut-offs, the sample floor and the ladder."""
    ws = wb[SCORECARD_SHEET]
    out = {}
    for r in range(1, ws.max_row + 1):
        label = _text(ws, r, 11).lower().rstrip(":").strip()
        value = _num(_cell(ws, r, 12))
        if value is None or not label:
            continue
        if label in CONSTANT_LABELS:
            out[CONSTANT_LABELS[label]] = value
        elif label in RATING_LABELS:
            out[f"rating_{label.replace(' ', '_')}"] = value
    missing = [k for k in ("score_excellent", "score_good", "score_fair",
                           "score_poor") if k not in out]
    if missing:
        raise WorkbookError(
            f"{SCORECARD_SHEET}: could not read the band scoring key "
            f"({', '.join(missing)}). Without it no band can be turned into "
            f"a score.")
    return out


# ---------------------------------------------------------------------------
# 6. Data Dictionary — extraction guidance, per parameter
# ---------------------------------------------------------------------------

def read_dictionary(wb):
    """What the sheet tells an analyst about each parameter.

    Keyed by ref code because the dictionary is written per assessment row,
    and several rows share an input key with their reference sibling.
    """
    if DICT_SHEET not in wb.sheetnames:
        return {}
    ws = wb[DICT_SHEET]
    out = {}
    for r in range(1, ws.max_row + 1):
        ref = _text(ws, r, 2)
        if not ref or ref.lower() in ("param", "category"):
            continue
        out[ref] = {
            "scoring_method": _text(ws, r, 5),
            "keys": _text(ws, r, 6),
            "definition": _text(ws, r, 7),
            "unit": _text(ws, r, 8),
            "where_to_find": _text(ws, r, 9),
            "if_missing": _text(ws, r, 10),
        }
    return out


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def parse_config_workbook(path):
    """Read every configuration sheet. Returns one dict, raises on structure.

    Deliberately strict: a workbook whose category weights do not sum to 100,
    or whose input sheet has a duplicate key, is not a workbook we can score
    against, and importing it half-read would be worse than refusing it.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=False)
    # A second, formula-preserving read. Cut-points and weights are values;
    # which benchmark column a percentile node reads is only in its formula.
    formulas = load_workbook(path, data_only=False, read_only=False)
    try:
        inputs = read_inputs(wb)
        monotonic, ranges = read_rubrics(wb)
        anchors = read_anchors(wb)
        categories, subitems, children, leaves = read_tree(wb, formulas)
        constants = read_constants(wb)
        dictionary = read_dictionary(wb)
    finally:
        wb.close()
        formulas.close()

    anchor_keys = {a["input_key"] for a in anchors}
    rubric_keys = ({m["input_key"] for m in monotonic}
                   | {r["input_key"] for r in ranges})
    scored_keys = {lf["input_key"] for lf in leaves if lf.get("input_key")}

    # A key on a category sheet that the input sheet does not collect can
    # never be answered — the same class of defect as a weighted sub-item
    # with nothing beneath it, one level down.
    declared = {i["input_key"] for i in inputs}
    orphan_leaves = sorted(scored_keys - declared)

    # A key the sheet collects but nothing scores. Reference rows are
    # legitimately in this set; anything else is a gap.
    unscored = sorted(declared - scored_keys)

    report = {
        "inputs": len(inputs),
        "monotonic_rubrics": len(monotonic),
        "range_rubrics": len(ranges),
        "anchors": len(anchors),
        "categories": len(categories),
        "subitems": len(subitems),
        "children": len(children),
        "scored_leaves": len(leaves),
        "percentile_nodes": len([lf for lf in leaves
                                 if lf.get("benchmark_metric")]),
        "dictionary_rows": len(dictionary),
        "orphan_leaves": orphan_leaves,
        "unscored_inputs": unscored,
        "constants": {k: float(v) for k, v in constants.items()},
    }
    if orphan_leaves:
        raise WorkbookError(
            f"these keys are scored on a category sheet but not collected on "
            f"the input sheet: {', '.join(orphan_leaves)}. They would score "
            f"blank for every company.")

    return {
        "inputs": inputs, "monotonic": monotonic, "ranges": ranges,
        "anchors": anchors, "categories": categories, "subitems": subitems,
        "children": children, "leaves": leaves, "constants": constants,
        "dictionary": dictionary, "anchor_keys": anchor_keys,
        "rubric_keys": rubric_keys, "report": report,
    }
